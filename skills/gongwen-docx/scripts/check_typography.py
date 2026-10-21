"""Check editable text against this user's government-report typography rules."""
import argparse
import json
import re
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
from docx.oxml import parse_xml
from docx.text.run import Run
from docx.text.paragraph import Paragraph
from docx import Document
from docx.oxml.ns import qn

FANG = {'FangSong_GB2312', '仿宋_GB2312'}
KAI = {'KaiTi_GB2312', '楷体_GB2312'}
HEI = {'Heiti SC', '黑体', 'SimHei'}
SONG = {'Songti SC', '宋体', 'SimSun'}
TITLE = {'FZXiaoBiaoSong-B05S', '方正小标宋简体', '方正小标宋简'}

CHINESE = r'\u3400-\u4dbf\u4e00-\u9fff'
SERIAL_SPACING = re.compile(
    r'^(?:[一二三四五六七八九十百千万零〇]+、|[（(][一二三四五六七八九十百千万零〇0-9]+[）)]|[0-9]+[.．、])'
    rf'[^\S\r\n]+(?=[{CHINESE}])'
)
INLINE_SERIAL = re.compile(rf'^\d+\.[^\S\r\n]*(?=[{CHINESE}])')

CHAPTER_SECTION = re.compile(r'^(第[一二三四五六七八九十百千万零〇0-9]+[章节])([^\S\r\n]*)(\S[^\r\n]*)$')

def serial_spacing_error(text, *, chapter_heading=False):
    if chapter_heading:
        match = CHAPTER_SECTION.match(text)
        if match and match.group(2) != ' ':
            return '章、节序号后须保留一个半角空格（U+0020）'
    if SERIAL_SPACING.match(text):
        return '其他层级序号后应直接接中文，不留空格或制表符'
    return None

def check_spacing(path):
    from docx.text.paragraph import Paragraph
    d = Document(path)
    errors = []
    counts = {'paragraphs': 0, 'chapter_section_headings': 0}
    for i, element in enumerate(d.element.xpath('.//w:p')):
        p = Paragraph(element, d._body)
        text = ''.join(element.xpath('.//w:t/text()'))
        if not text.strip():
            continue
        counts['paragraphs'] += 1
        own = element.find('w:pPr/w:outlineLvl', element.nsmap)
        inherited = p.style.element.find('w:pPr/w:outlineLvl', p.style.element.nsmap)
        outline = own if own is not None else inherited
        is_heading = (outline is not None and int(outline.get(qn('w:val'))) < 9) or p.style.name.startswith('Heading ')
        in_toc = any(a.tag == qn('w:sdt') and any(x.get(qn('w:val')) == 'Table of Contents' for x in a.iter(qn('w:docPartGallery'))) for a in element.iterancestors())
        chapter_heading = (is_heading or in_toc) and bool(CHAPTER_SECTION.match(text))
        counts['chapter_section_headings'] += int(chapter_heading)
        issue = serial_spacing_error(text, chapter_heading=chapter_heading)
        if issue:
            errors.append(f'段落{i}: {issue}：{text[:60]}')
    return {'path': str(Path(path).resolve()), 'scope': 'serial_spacing', 'passed': not errors, 'counts': counts, 'errors': errors}

def paragraph_property(p, tag, attr='val', default=None):
    containers = [p._p.pPr]
    style, seen = p.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        containers.append(style.element.pPr)
        style = style.base_style
    for container in containers:
        el = None if container is None else container.find(qn('w:' + tag))
        if el is not None:
            value = el.get(qn('w:' + attr))
            if value is not None:
                return value
            if attr == 'val':
                return '1'
    return default


def all_runs(p):
    # Hyperlinks and cached fields are intentionally included.
    return [Run(el, p) for el in p._p.xpath('.//w:r')]


def check(path, allow_heading_line_breaks=False, body_line_pt=None):
    d = Document(path)
    errors = []
    counts = {'runs': 0, 'headings': 0, 'inline_titles': 0, 'captions': 0, 'tables': len(d.tables), 'heading_manual_breaks': 0, 'title_manual_breaks': 0}

    lines = Counter()
    for p in d.paragraphs:
        if p.text.strip() and paragraph_property(p, 'outlineLvl', default='9') == '9' and paragraph_property(p, 'spacing', 'lineRule') == 'exact' and p.style.name not in ('Caption', 'GovCaption', 'GovNote', 'GovSource', 'GovLegend', 'GovAxis', 'GovDate'):
            value = paragraph_property(p, 'spacing', 'line')
            if value:
                lines[int(value)] += 1
    line_twips = round(body_line_pt * 20) if body_line_pt is not None else (lines.most_common(1)[0][0] if lines else None)
    if line_twips is None:
        errors.append('未识别固定正文行距，请使用--body-line-pt明确指定')

    def run_check(run, families, bold, size, where):
        if not run.text:
            return
        counts['runs'] += 1
        citation = run.style.name in ('GovCitation', 'Citation', 'Footnote Reference', 'Endnote Reference') or (run.font.superscript and bool(re.fullmatch(r'[\[［]?[0-9,，、;；—–\- ]+[\]］]?', run.text)))
        if citation:
            families, bold, size = FANG, False, 12
        rp = run._r.rPr
        rf = None if rp is None else rp.find(qn('w:rFonts'))
        for channel in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            value = None if rf is None else rf.get(qn('w:' + channel))
            if value not in families:
                errors.append(f'{where}: {channel}字体为{value}')
        if rf is not None and any(k.endswith('Theme') or k.endswith('theme') for k in rf.attrib):
            errors.append(f'{where}: 存在主题字体覆盖')
        for tag in ('b', 'bCs'):
            el = None if rp is None else rp.find(qn('w:' + tag))
            value = None if el is None else el.get(qn('w:val'), '1')
            actual = value in ('1', 'true', 'on') if value is not None else None
            if actual != bold:
                errors.append(f'{where}: {tag}={value}，应为{int(bold)}')
        if size is not None and (run.font.size is None or abs(run.font.size.pt - size) > .01):
            errors.append(f'{where}: 字号不符{size}磅')

    for i, p in enumerate(d.paragraphs):
        text = ''.join(p._p.xpath('.//w:t/text()'))
        if not text.strip():
            continue
        name = p.style.name
        issue = serial_spacing_error(text, chapter_heading=name.startswith('Heading ') or name in ('GovChapter', 'GovSection'))
        if issue:
            errors.append(f'段落{i}: {issue}')
        is_chapter = bool(re.match(r'^第[一二三四五六七八九十百千万零〇0-9]+章', text)) and (name in ('GovChapter', 'GovAppendix') or int(paragraph_property(p, 'outlineLvl', default='9')) < 9 or name.startswith('Heading '))
        if name == 'Title':
            counts['title_manual_breaks'] += len(p._p.xpath('.//w:br[not(@w:type) or @w:type="textWrapping"] | .//w:cr'))
        is_heading = is_chapter or int(paragraph_property(p, 'outlineLvl', default='9')) < 9 or name in ('GovSection', 'GovSubsection', 'GovClose', 'GovAppendix') or name.startswith('Heading ') or bool(re.match(r'^第[一二三四五六七八九十百0-9]+节', text))
        if is_heading:
            breaks = p._p.xpath('.//w:br[not(@w:type) or @w:type="textWrapping"] | .//w:cr')
            counts['heading_manual_breaks'] += len(breaks)
            if breaks and not allow_heading_line_breaks:
                errors.append(f'段落{i}: 层级标题含{len(breaks)}处手工断行，应按版心自然换行；仅用户明确指定时可保留')
        if name == 'Title' or is_chapter:
            family, bold, size = TITLE, False, 22
        elif name in ('Heading 1', 'GovSection', 'GovPoint', 'GovClose', 'GovAppendix') or re.match(r'^第[一二三四五六七八九十百0-9]+节', text):
            family, bold, size = HEI, True, 16
            counts['headings'] += 1
        elif name in ('Heading 2', 'GovSubsection'):
            family, bold, size = KAI, True, 16
            counts['headings'] += 1
        elif name in ('Heading 3', 'Heading 4'):
            family, bold, size = FANG, True, 16
            counts['headings'] += 1
        elif name in ('Caption', 'GovCaption') or re.match(r'^[图表]\d+：', text):
            family, bold, size = FANG, True, 14
            counts['captions'] += 1
            if not re.match(r'^[图表]\d+：\S',text):
                errors.append(f'段落{i}: 图表编号后须紧接全角冒号和标题')
        elif INLINE_SERIAL.match(text) and '。' in text:
            counts['inline_titles'] += 1
            limit = text.index('。') + 1
            offset = 0
            for run in all_runs(p):
                end = offset + len(run.text)
                if offset < limit < end:
                    errors.append(f'段落{i}: 段内标题和解释未分开设置字重')
                run_check(run, FANG, offset < limit, 16, f'段落{i}')
                offset = end
            continue
        elif name in ('GovNote', 'GovSource', 'GovLegend', 'GovAxis', 'GovDate', 'Footnote Text', 'Endnote Text') or text.startswith(('注：', '资料来源：', '数据来源：')):
            family, bold, size = FANG, False, 12
        else:
            family, bold, size = FANG, False, 16
        if name == 'Title' or is_heading or name == 'GovAppendix':
            style = p.style
            seen = set()
            containers = [p._p.pPr]
            while style is not None and style.style_id not in seen:
                seen.add(style.style_id)
                containers.append(style.element.pPr)
                style = style.base_style
            if any(c is not None and (c.find(qn('w:pBdr')) is not None or c.find(qn('w:shd')) is not None) for c in containers):
                errors.append(f'段落{i}: 标题存在段落边框或底纹，须核对并清除无依据的模板装饰')
        for run in all_runs(p):
            run_check(run, family, bold, size, f'段落{i}')

    # Group caption, drawing/table, legend and source notes as one block.
    body = [e for e in d.element.body if e.tag in (qn('w:p'), qn('w:tbl'))]
    def para(el):
        return Paragraph(el, d._body) if el.tag == qn('w:p') else None
    def note(el):
        p = para(el)
        return p is not None and (p.style.name in ('GovNote', 'GovSource', 'GovLegend', 'GovAxis') or p.text.startswith(('注：', '资料来源：', '数据来源：')))
    def caption(el):
        p = para(el)
        return p is not None and (p.style.name == 'Caption' or bool(re.match(r'^[图表]\d+：', p.text)))
    consumed = set()
    for i, el in enumerate(body):
        if i in consumed:
            continue
        visual = el.tag == qn('w:tbl') or bool(el.xpath('.//w:drawing'))
        if not (caption(el) or visual):
            continue
        start, end, has_visual = i, i, visual
        if caption(el):
            if i + 1 < len(body) and (body[i+1].tag == qn('w:tbl') or body[i+1].xpath('.//w:drawing')):
                end, has_visual = i+1, True
            else:
                errors.append(f'图表标题{i}: 后面未紧接所属图表')
        while end + 1 < len(body) and note(body[end+1]):
            end += 1
        consumed.update(range(start, end+1))
        if not has_visual or line_twips is None:
            continue
        before = para(body[start])
        prior = para(body[start-1]) if start else None
        # Adjacent paragraphs use the larger space in normal Word layout;
        # page-top suppression and table-edge behavior still require rendering.
        gap_before = max(int(paragraph_property(before, 'spacing', 'before', '0')) if before else 0,
                         int(paragraph_property(prior, 'spacing', 'after', '0')) if prior else 0)
        if start and abs(gap_before-line_twips)>2:
            errors.append(f'图表整体{start}: 前方间距{gap_before} twips，应为正文一行{line_twips}')
        last = para(body[end])
        following = para(body[end+1]) if end+1<len(body) else None
        gap_after = max(int(paragraph_property(last, 'spacing', 'after', '0')) if last else 0,
                        int(paragraph_property(following, 'spacing', 'before', '0')) if following else 0)
        if end+1<len(body) and abs(gap_after-line_twips)>2:
            errors.append(f'图表整体{start}: 后方间距{gap_after} twips，应为正文一行{line_twips}')

    for ti, table in enumerate(d.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    for run in all_runs(p):
                        run_check(run, FANG, ri == 0, 12, f'表{ti+1}行{ri+1}列{ci+1}')

    for si, section in enumerate(d.sections):
        for prop, expected in [('page_width',210),('page_height',297),('top_margin',37),('bottom_margin',35),('left_margin',28),('right_margin',26)]:
            value = getattr(section, prop).mm
            if abs(value - expected) > .1:
                errors.append(f'节{si+1}: {prop}={value:.2f}毫米')
        for label, footer in [('奇数',section.footer),('偶数',section.even_page_footer)]:
            # Include cached field results, which Paragraph.runs does not expose.
            from docx.text.run import Run
            for p in footer.paragraphs:
                for element in p._p.xpath('.//w:r'):
                    run_check(Run(element, p), SONG, False, 14, f'{label}页码')
    # Footnote/endnote parts do not occur in Document.paragraphs.
    with ZipFile(path) as package:
        for part in ('word/footnotes.xml', 'word/endnotes.xml'):
            if part not in package.namelist():
                continue
            root = parse_xml(package.read(part))
            for pi, pe in enumerate(root.iter(qn('w:p'))):
                p = Paragraph(pe, d._body)
                for run in all_runs(p):
                    run_check(run, FANG, False, 12, f'{part}段落{pi}')
    return {'path': str(Path(path).resolve()), 'scope': 'editable_ooxml_typography', 'passed': not errors, 'body_line_twips': line_twips, 'counts': counts, 'errors': errors, 'render_review_required': True}

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('docx')
    parser.add_argument('--output-json')
    parser.add_argument('--body-line-pt', type=float, help='正文固定行距；省略时从正文识别')
    parser.add_argument('--allow-heading-line-breaks', action='store_true', help='仅在用户明确指定标题断行位置或要求保留已有标题版式时使用')
    parser.add_argument('--spacing-only', action='store_true', help='仅检查章、节及其他层级序号后的间隔，包含目录')
    args = parser.parse_args()
    report = check_spacing(args.docx) if args.spacing_only else check(args.docx, allow_heading_line_breaks=args.allow_heading_line_breaks, body_line_pt=args.body_line_pt)
    output = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output_json:
        Path(args.output_json).write_text(output)
    print(output)
    raise SystemExit(0 if report['passed'] else 1)
