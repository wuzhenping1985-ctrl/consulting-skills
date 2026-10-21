"""Local CIIC brand enforcement, added 2026-09-10.

Preserves editable OOXML. Never remaps arbitrary user colors to nearest colors.
Photographs and embedded assets require visual review and are not recolored.
"""
import re
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL_TYPE
from .constants import BRAND, ALLOWED_HEX, THEME_COLORS, FONT_BODY, CHART_COLORS

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS = {'a': A}
FONT_TAGS = ('latin', 'ea', 'cs')
ALIASES = {'bg1':'lt1', 'tx1':'dk1', 'bg2':'lt2', 'tx2':'dk2'}


def _font_props(props):
    # DrawingML schema requires latin, ea, cs in this order, before hyperlinks.
    for tag in FONT_TAGS:
        for child in list(props.findall(f'{{{A}}}{tag}')):
            props.remove(child)
    before = next((x for x in props if etree.QName(x).localname in ('sym','hlinkClick','hlinkMouseOver','rtl','extLst')), None)
    index = props.index(before) if before is not None else len(props)
    for offset, tag in enumerate(FONT_TAGS):
        child = etree.Element(f'{{{A}}}{tag}', typeface=FONT_BODY)
        props.insert(index + offset, child)


def normalize_xml(root, filename):
    """Set complete font coverage and exact theme colors during export."""
    if not filename.startswith('ppt/'):
        return
    for el in list(root.iter()):
        if 'typeface' in el.attrib:
            el.set('typeface', FONT_BODY)
        if el.tag in (f'{{{A}}}rPr', f'{{{A}}}defRPr', f'{{{A}}}endParaRPr'):
            _font_props(el)
    # Explicit font on every existing text run, including fields and chart labels.
    for text in root.findall(f'.//{{{A}}}t'):
        parent = text.getparent()
        props = parent.find(f'{{{A}}}rPr')
        if props is None:
            props = etree.Element(f'{{{A}}}rPr')
            parent.insert(0, props)
        _font_props(props)
    if filename.startswith('ppt/theme/'):
        for scheme in root.findall(f'.//{{{A}}}clrScheme'):
            scheme.set('name', 'CIIC')
            for key, value in THEME_COLORS.items():
                node = scheme.find(f'{{{A}}}{key}')
                if node is None:
                    node = etree.SubElement(scheme, f'{{{A}}}{key}')
                for child in list(node):
                    node.remove(child)
                etree.SubElement(node, f'{{{A}}}srgbClr', val=value)
        for scheme in root.findall(f'.//{{{A}}}fontScheme'):
            scheme.set('name', 'Microsoft YaHei')
        # The stock Office effects theme uses black for shadows/lines.
        for color in root.findall(f'.//{{{A}}}srgbClr'):
            if color.get('val', '').upper() == '000000':
                color.set('val', BRAND['text'])


def _resolve(el):
    local = etree.QName(el).localname
    if local == 'srgbClr':
        value = el.get('val', '').upper()
    elif local == 'schemeClr':
        key = ALIASES.get(el.get('val'), el.get('val'))
        value = THEME_COLORS.get(key)
    elif local == 'sysClr':
        value = el.get('lastClr', '').upper()
    else:
        return None
    # Arbitrary modifiers change the specified palette. Export uses solid colors.
    if len(el):
        return None
    return value


def contrast(a, b):
    def luminance(value):
        components = [int(value[i:i+2], 16)/255 for i in (0, 2, 4)]
        components = [c/12.92 if c <= 0.04045 else ((c+0.055)/1.055)**2.4 for c in components]
        return sum(c*w for c,w in zip(components, (.2126,.7152,.0722)))
    x,y=sorted((luminance(a),luminance(b)))
    return (y+.05)/(x+.05)


def _solid(fill):
    try:
        if fill.type != MSO_FILL_TYPE.SOLID:
            return None
        rgb=fill.fore_color.rgb
        if rgb is not None:
            return str(rgb)
    except (AttributeError, TypeError):
        pass
    try:
        holder=fill._xPr.find(f'{{{A}}}solidFill')
        return _resolve(holder[0]) if holder is not None and len(holder) else None
    except (AttributeError, IndexError):
        return None


def _contains(back, front):
    tol=9144  # .01 inch avoids round-off at matching box edges.
    return (back.left <= front.left+tol and back.top <= front.top+tol and
            back.left+back.width >= front.left+front.width-tol and
            back.top+back.height >= front.top+front.height-tol)


def _text_boxes(prs):
    """Yield native text frames with solid backgrounds resolved in drawing order.

    Complex intersections and pictures remain a required visual review step.
    """
    for n,slide in enumerate(prs.slides,1):
        base=_solid(slide.background.fill) or BRAND['white']
        def walk(shapes, background):
            preceding=[]
            for shape in shapes:
                bg=background
                for back in reversed(preceding):
                    if _contains(back,shape):
                        candidate=_solid(back.fill) if hasattr(back,'fill') else None
                        if candidate:
                            bg=candidate
                            break
                own=_solid(shape.fill) if hasattr(shape,'fill') else None
                bg=own or bg
                if shape.has_text_frame:
                    yield n,shape.name,shape.text_frame,bg
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            yield n,shape.name,cell.text_frame,_solid(cell.fill) or bg
                if hasattr(shape,'shapes'):
                    yield from walk(shape.shapes,bg)
                preceding.append(shape)
        yield from walk(slide.shapes,base)


def _text_color(run, paragraph):
    for props in (run._r.find(f'{{{A}}}rPr'), paragraph._p.find(f'{{{A}}}pPr/{{{A}}}defRPr')):
        if props is not None:
            fill=props.find(f'{{{A}}}solidFill')
            if fill is not None and len(fill):
                return _resolve(fill[0])
    return BRAND['text']


def _minimum(run, paragraph):
    size = run.font.size or paragraph.font.size
    bold = run.font.bold if run.font.bold is not None else paragraph.font.bold
    large = size is not None and (size.pt >= 18 or (bool(bold) and size.pt >= 14))
    return 3.0 if large else BRAND['contrast_minimum']


def improve_text_contrast(prs):
    """Use palette white/dark-gray when current palette text lacks contrast."""
    # Explicit fills prevent stock table styles from introducing derived tints.
    def prepare_tables(shapes):
        for shape in shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.fill.type is None:
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor.from_string(BRAND['white'])
            if hasattr(shape, 'shapes'):
                prepare_tables(shape.shapes)
    for slide in prs.slides:
        prepare_tables(slide.shapes)
    for _,_,frame,bg in _text_boxes(prs):
        for p in frame.paragraphs:
            for run in p.runs:
                if not run.text.strip():
                    continue
                fg=_text_color(run,p)
                # Unknown colors must fail validation, never silently disappear.
                if fg not in ALLOWED_HEX or bg not in ALLOWED_HEX:
                    continue
                if contrast(fg,bg) < _minimum(run,p):
                    candidates=(BRAND['white'],BRAND['text'])
                    fg=max(candidates,key=lambda value:contrast(value,bg))
                run.font.color.rgb=RGBColor.from_string(fg)
    # Native data charts also need an explicit brand series order.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            for i,series in enumerate(shape.chart.series):
                if series.format.fill.type is None:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb=CHART_COLORS[i%len(CHART_COLORS)]


def inspect_brand(path, check_contrast=True):
    """Independent read-only final-file check, also used by the S4 gate."""
    errors=[]
    def error(part,category,message):
        m=re.search(r'slides/slide(\d+)\.xml',part)
        errors.append({'slide':int(m.group(1)) if m else 0,'category':category,'message':message,'part':part})
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.startswith('ppt/') or not name.endswith('.xml'):
                continue
            root=etree.fromstring(z.read(name))
            for el in root.iter():
                if 'typeface' in el.attrib and el.get('typeface') != FONT_BODY:
                    error(name,'brand_font',f"Font must be {FONT_BODY}: {el.get('typeface')}")
            for text in root.findall(f'.//{{{A}}}t'):
                props=text.getparent().find(f'{{{A}}}rPr')
                if props is None or any(props.find(f'{{{A}}}{tag}') is None or props.find(f'{{{A}}}{tag}').get('typeface') != FONT_BODY for tag in FONT_TAGS):
                    error(name,'brand_font','Text run missing explicit Latin/East Asian/complex-script font')
            if name.startswith('ppt/theme/'):
                for scheme in root.findall(f'.//{{{A}}}clrScheme'):
                    for key,value in THEME_COLORS.items():
                        node=scheme.find(f'{{{A}}}{key}/{{{A}}}srgbClr')
                        if node is None or node.get('val','').upper()!=value:
                            error(name,'brand_theme',f'{key} must be #{value}')
            if not name.startswith(('ppt/slides/','ppt/charts/','ppt/diagrams/')):
                continue
            for el in root.iter():
                if etree.QName(el).namespace != A:
                    continue
                local=etree.QName(el).localname
                if local in ('srgbClr','schemeClr','sysClr','prstClr','scrgbClr','hslClr'):
                    value=_resolve(el)
                    if value not in ALLOWED_HEX:
                        error(name,'brand_color',f"Non-palette or modified color: {local} {dict(el.attrib)}")
                elif local in ('gradFill','pattFill'):
                    error(name,'brand_color','Native gradients/patterns need replacement with prescribed solid palette')
    if check_contrast:
        for n,shape,frame,bg in _text_boxes(Presentation(path)):
            for p in frame.paragraphs:
                for run in p.runs:
                    if not run.text.strip():
                        continue
                    fg=_text_color(run,p)
                    if fg in ALLOWED_HEX and bg in ALLOWED_HEX and contrast(fg,bg)<_minimum(run,p):
                        errors.append({'slide':n,'category':'brand_contrast','shape':shape,'message':f'{fg} on {bg}: {contrast(fg,bg):.2f}:1; {run.text[:50]}'})
    return {'passed':not errors,'errors':errors,'font':FONT_BODY,'palette':sorted(ALLOWED_HEX)}
