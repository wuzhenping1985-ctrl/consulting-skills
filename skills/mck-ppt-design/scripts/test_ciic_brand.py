"""Regression tests for the local CIIC customization; no network or app writes."""
import importlib.util
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from mck_ppt import MckEngine
from mck_ppt.constants import *
from mck_ppt.core import add_text, add_rect
from mck_ppt.brand import inspect_brand
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

spec=importlib.util.spec_from_file_location('ciic_gate',ROOT/'references/scripts/gate_check.py')
gate=importlib.util.module_from_spec(spec)
spec.loader.exec_module(gate)
A='http://schemas.openxmlformats.org/drawingml/2006/main'


def fixture(path):
    eng=MckEngine(total_slides=6)
    eng.cover(title='中智咨询字体与配色',subtitle='中文、English 与数字 2026',date='2026年9月')
    eng.four_column(title='字体、配色、图表与检查采用统一规范',items=[('1','字体','中文、英文与数字\n统一使用微软雅黑'),('2','配色','主蓝色与灰色\n明确区分层级'),('3','图表','按规定顺序用色\n保留标签和图例'),('4','检查','验证原生对象\n逐页检查可读性')])
    eng.grouped_bar(title='多系列数据使用规定的蓝色顺序',categories=['方案甲','方案乙','方案丙'],series=[('系列一',CHART_COLORS[0]),('系列二',CHART_COLORS[1])],data=[[120,90],[100,110],[130,100]],max_val=150,y_ticks=[0,50,100,150],source='注：仅用于样式验证的虚拟数据')
    eng.matrix_2x2(title='浅底深字与蓝底白字保持清晰层次',quadrants=[('主蓝底',PRIMARY_BLUE,'该区域自动采用白色文字'),('浅蓝底',CHART_COLORS[3],'该区域采用深灰色文字'),('浅灰底',BG_GRAY,'该区域采用深灰色文字'),('浅蓝底',CHART_COLORS[4],'该区域采用深灰色文字')])
    s=eng._ns()
    add_text(s,LM,Inches(.3),CW,Inches(.5),'主题色板与原生图表',font_size=ACTION_TITLE_SIZE,font_color=PRIMARY_BLUE)
    for i,c in enumerate(CHART_COLORS):
        add_rect(s,LM+Inches(1.8)*i,Inches(1.1),Inches(1.5),Inches(.6),c)
        add_text(s,LM+Inches(1.8)*i,Inches(1.8),Inches(1.5),Inches(.3),f'{i+1}: #{c}',font_size=Pt(11))
    data=CategoryChartData(); data.categories=['甲','乙','丙']
    data.add_series('样本一',[10,20,30]);data.add_series('样本二',[15,25,20])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,LM,Inches(2.6),CW,Inches(3.7),data)
    s.notes_slide.notes_text_frame.text='备注字体验证：Microsoft YaHei 中文 123'
    s=eng._ns()
    add_text(s,LM,Inches(.3),CW,Inches(.5),'原生表格和文本对比度',font_size=ACTION_TITLE_SIZE,font_color=PRIMARY_BLUE)
    table=s.shapes.add_table(2,2,LM,Inches(1.2),CW,Inches(2)).table
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells):
            cell.text=f'单元格 {i+1}-{j+1} ABC 123'
            cell.fill.solid();cell.fill.fore_color.rgb=BG_GRAY
            for p in cell.text_frame.paragraphs:
                p.font.name='Arial';p.font.size=Pt(14)
                for run in p.runs:run.font.color.rgb=WHITE
    add_rect(s,LM,Inches(4),CW,Inches(1),BG_GRAY)
    add_text(s,LM,Inches(4.2),CW,Inches(.5),'浅灰底白字会自动修正为深灰字',font_size=BODY_SIZE,font_color=WHITE)
    eng.save(str(path))


class BrandTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(); cls.base=Path(cls.tmp.name)/'sample.pptx'
        fixture(cls.base)
    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_export_is_editable_and_branded(self):
        self.assertTrue(inspect_brand(self.base)['passed'])
        prs=Presentation(self.base)
        self.assertEqual(len(prs.slides),6)
        self.assertTrue(any(s.has_chart for slide in prs.slides for s in slide.shapes))
        self.assertTrue(any(s.has_table for slide in prs.slides for s in slide.shapes))
        with zipfile.ZipFile(self.base) as z:
            fonts=set()
            for name in z.namelist():
                if name.startswith('ppt/') and name.endswith('.xml'):
                    fonts.update(etree.fromstring(z.read(name)).xpath('//@typeface'))
            self.assertEqual(fonts,{'Microsoft YaHei'})

    def test_off_palette_generation_fails(self):
        eng=MckEngine(total_slides=1);s=eng._ns()
        add_rect(s,LM,Inches(1),CW,Inches(2),RGBColor(0,255,0))
        with self.assertRaisesRegex(ValueError,'CIIC brand validation failed'):
            eng.save(str(Path(self.tmp.name)/'green.pptx'))

    def test_gate_rejects_tampered_font_palette_and_contrast(self):
        for category in ('brand_font','brand_color','brand_contrast'):
            bad=Path(self.tmp.name)/f'{category}.pptx'
            with zipfile.ZipFile(self.base) as src,zipfile.ZipFile(bad,'w') as dst:
                for item in src.infolist():
                    blob=src.read(item.filename)
                    if item.filename=='ppt/slides/slide1.xml':
                        root=etree.fromstring(blob)
                        if category=='brand_font':
                            root.find(f'.//{{{A}}}latin').set('typeface','Arial')
                        elif category=='brand_color':
                            root.find(f'.//{{{A}}}srgbClr').set('val','00FF00')
                        else:
                            text=root.find(f'.//{{{A}}}t')
                            props=text.getparent().find(f'{{{A}}}rPr')
                            for old in list(props.findall(f'{{{A}}}solidFill')):
                                props.remove(old)
                            fill=etree.Element(f'{{{A}}}solidFill')
                            props.insert(0,fill)
                            etree.SubElement(fill,f'{{{A}}}srgbClr',val='E6E6E6')
                        blob=etree.tostring(root)
                    dst.writestr(item,blob)
            result=gate.run_gate_check(str(bad),self.tmp.name)
            self.assertFalse(result['passed'],category)
            self.assertIn(category,{e['category'] for e in result['user_code_error_detail']})

if __name__=='__main__':
    if len(sys.argv)==3 and sys.argv[1]=='--sample':
        fixture(Path(sys.argv[2]))
    else:unittest.main()
